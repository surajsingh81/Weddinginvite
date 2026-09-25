#!/usr/bin/env python3
"""
Create wedding-details.xlsx - the editable master sheet for the wedding website.

Run:  python3 make_workbook.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "wedding-details.xlsx")

FILLIN = "-- FILL IN --"

MAROON = "7B1E2B"
GOLD = "C9A227"
CREAM = "FDF6E9"
GREY = "EFEFEF"

head_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor=MAROON)
title_font = Font(name="Calibri", size=11, bold=True, color=MAROON)
fill_font = Font(name="Calibri", size=11, color="B8860B", bold=True)
fill_fill = PatternFill("solid", fgColor=CREAM)
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ---------------------------------------------------------------- Wedding Details
ws = wb.active
ws.title = "Wedding Details"
rows = [
    ("Field", "Value", "Notes"),
    ("Groom Name", FILLIN, "Decorative font on card - not machine readable"),
    ("Bride Name", FILLIN, "Decorative font on card - not machine readable"),
    ("Groom's Parents", FILLIN, "As printed on the card"),
    ("Bride's Parents", FILLIN, "As printed on the card"),
    ("Wedding Year", "2026", "Year is not printed on the card; used as fallback"),
    ("Wedding Main Event", "Vidai (Milap)", "Read from card"),
    ("Wedding Date", "12 Dec 2026", "Read from card: 12 Dec, 4:00 AM"),
    ("Wedding Start Time", "04:00 AM", "Read from card"),
    ("Function Dates", "09 Dec 2026 to 12 Dec 2026", "Read from card"),
    ("City", FILLIN, "Not machine readable from card"),
    ("Primary Venue Name", FILLIN, "Wedding / main function venue"),
    ("Primary Venue Address", FILLIN, "Full address for the website"),
    ("Google Maps Link", FILLIN, "Paste a maps.app.goo.gl or google.com/maps link"),
    ("Other Venue 1 Name", FILLIN, "Optional: haldi or baraat venue"),
    ("Other Venue 1 Address", FILLIN, "Optional"),
    ("Other Venue 2 Name", FILLIN, "Optional"),
    ("Other Venue 2 Address", FILLIN, "Optional"),
    ("Contact Name", FILLIN, "Who answers RSVP calls"),
    ("Contact Phone", FILLIN, "With country code, e.g. +91 98765 43210"),
    ("WhatsApp Number", FILLIN, "With country code; powers the RSVP button"),
    ("Email", FILLIN, "Optional"),
    ("RSVP Deadline", FILLIN, "e.g. 01 Dec 2026"),
    ("Website Title", "We're Getting Married", "Browser tab and hero heading"),
    ("Couple Line (short)", FILLIN, "e.g. Aarav weds Aisha"),
    ("Opening Blessing Line", "With the blessings of our families", "Above the couple's names"),
    ("Hashtag", FILLIN, "e.g. #AaravWedsAisha2026"),
    ("Quote / Verse", FILLIN, "Optional shloka shown near the end"),
    ("Accent Style", "classic", "classic | royal | temple - changes site colours"),
]
for r in rows:
    ws.append(list(r))

for c in ws[1]:
    c.font = head_font
    c.fill = head_fill
    c.border = border
    c.alignment = Alignment(horizontal="center", vertical="center")

for row in ws.iter_rows(min_row=2):
    row[0].font = title_font
    for c in row:
        c.border = border
        c.alignment = wrap
    v = row[1].value
    if isinstance(v, str) and v.strip() == FILLIN:
        row[1].font = fill_font
        row[1].fill = fill_fill
    elif v:
        row[1].font = Font(name="Calibri", size=11, bold=True, color="1F6F43")
        row[1].fill = PatternFill("solid", fgColor="EAF6EE")

ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 40
ws.column_dimensions["C"].width = 46
ws.freeze_panes = "A2"

dv = DataValidation(type="list", formula1='"classic,royal,temple"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"B{len(rows)}")

# ---------------------------------------------------------------- Functions
ws2 = wb.create_sheet("Wedding Functions")
funcs = [
    ("#", "Date", "Event", "Time", "Venue", "Note (shown on site)", "Verified?"),
    (1, "09 Dec 2026", "Tilak", "02:00 PM - 04:00 PM", FILLIN, "", "No"),
    (2, "09 Dec 2026", "Matkor", "04:00 PM - 04:30 PM", FILLIN, "", "No"),
    (3, "09 Dec 2026", "Madwa", "06:00 PM - 07:00 PM", FILLIN, "", "No"),
    (4, "10 Dec 2026", "Haldi", "10:00 PM - 12:00 AM", FILLIN, "Ends at midnight", "No"),
    (5, "10 Dec 2026", "Kalra / Devpuji", "12:00 PM - 01:00 PM", FILLIN, "May be one event - please confirm", "No"),
    (6, "10 Dec 2026", "Dhid Hari / Ghee Dhari", "03:00 PM - 08:00 PM", FILLIN, "Ladies function", "No"),
    (7, "11 Dec 2026", "Janau", "12:00 PM", FILLIN, "", "No"),
    (8, "11 Dec 2026", "Bhunga Lawa", "03:00 PM", FILLIN, "", "No"),
    (9, "11 Dec 2026", "Bardekhai", "05:00 PM - 07:00 PM", FILLIN, "", "No"),
    (10, "11 Dec 2026", "Parat", "03:00 PM", FILLIN, "", "No"),
    (11, "11 Dec 2026", "Darwagar", "07:00 PM", FILLIN, "", "No"),
    (12, "12 Dec 2026", "Vidai (Milap)", "04:00 AM", FILLIN, "Main ceremony", "No"),
]
for r in funcs:
    ws2.append(list(r))
for c in ws2[1]:
    c.font = head_font
    c.fill = head_fill
    c.border = border
    c.alignment = Alignment(horizontal="center", vertical="center")
for row in ws2.iter_rows(min_row=2):
    for c in row:
        c.border = border
    row[2].font = Font(name="Calibri", size=11, bold=True, color=MAROON)
    if isinstance(row[4].value, str) and row[4].value.strip() == FILLIN:
        row[4].font = fill_font
        row[4].fill = fill_fill
    # the Verified column is a checklist, so leave it looking editable
    row[6].font = Font(name="Calibri", size=11, bold=True, color="B8860B")
    row[6].fill = fill_fill
for col, w in zip("ABCDEFG", (6, 16, 26, 24, 26, 34, 12)):
    ws2.column_dimensions[col].width = w
ws2.freeze_panes = "A2"

dv_v = DataValidation(type="list", formula1='"No,Yes"', allow_blank=True)
ws2.add_data_validation(dv_v)
dv_v.add("G2:G200")

# ---------------------------------------------------------------- Venues
ws3 = wb.create_sheet("Venues")
vrows = [
    ("Name", "Address", "Google Maps Link", "Which functions"),
    (FILLIN, FILLIN, "", "Wedding / main function"),
    (FILLIN, FILLIN, "", "Functions held elsewhere"),
    (FILLIN, FILLIN, "", "Optional"),
]
for r in vrows:
    ws3.append(list(r))
for c in ws3[1]:
    c.font = head_font
    c.fill = head_fill
    c.border = border
    c.alignment = Alignment(horizontal="center", vertical="center")
for row in ws3.iter_rows(min_row=2):
    for c in row:
        c.border = border
        c.alignment = wrap
    for c in (row[0], row[1]):
        if isinstance(c.value, str) and c.value.strip() == FILLIN:
            c.font = fill_font
            c.fill = fill_fill
for col, w in zip("ABCD", (28, 46, 34, 30)):
    ws3.column_dimensions[col].width = w
ws3.freeze_panes = "A2"

# ---------------------------------------------------------------- RSVP
ws4 = wb.create_sheet("RSVP List")
rrows = [
    ("Guest Name", "Relation", "Side", "Phone", "RSVP", "Guests", "Meal"),
    ("-- example --", "Uncle", "Groom", "", "Pending", 2, "Veg"),
    ("-- example --", "Aunt", "Bride", "", "Attending", 3, "Jain"),
]
for r in rrows:
    ws4.append(list(r))
for c in ws4[1]:
    c.font = head_font
    c.fill = head_fill
    c.border = border
    c.alignment = Alignment(horizontal="center", vertical="center")
for row in ws4.iter_rows(min_row=2):
    for c in row:
        c.border = border
dv2 = DataValidation(type="list", formula1='"Attending,Pending,Declined"', allow_blank=True)
ws4.add_data_validation(dv2)
dv2.add("E2:E200")
dv3 = DataValidation(type="list", formula1='"Veg,Jain,Non-veg"', allow_blank=True)
ws4.add_data_validation(dv3)
dv3.add("G2:G200")
for col, w in zip("ABCDEFG", (24, 16, 12, 18, 14, 10, 12)):
    ws4.column_dimensions[col].width = w
ws4.freeze_panes = "A2"

# ---------------------------------------------------------------- Help
ws5 = wb.create_sheet("How To Use")
help_rows = [
    ("1. EDIT HERE", ""),
    ("Fill in every '-- FILL IN --' cell on the Wedding Details and Venues sheets.", ""),
    ("", ""),
    ("2. SYNC TO WEBSITE", ""),
    ("Save this file, then run:", "python3 build_site.py"),
    ("This rewrites site/data.js from this workbook, and the website picks up your edits.", ""),
    ("", ""),
    ("3. WHAT CAME FROM THE PHOTO", ""),
    ("Read automatically: the 12 functions with their dates and times, spanning 09 Dec to 12 Dec.", ""),
    ("Not machine readable, please type them: names, parents, venue, address, city, year.", ""),
    ("", ""),
    ("4. DATE AND TIME FORMAT", ""),
    ("Dates: '09 Dec 2026'. Leave the year off and the Wedding Year field is used instead.", ""),
    ("Times: '04:00 PM' or '04:00 PM - 06:00 PM'. 24-hour also works.", ""),
    ("", ""),
    ("5. LIVE EDITS IN THE BROWSER", ""),
    ("The website also has an Edit button: type directly on the page, then Download JSON.", ""),
    ("That JSON has the same shape as this workbook, so either one can drive the site.", ""),
]
for r in help_rows:
    ws5.append(list(r))
for row in ws5.iter_rows():
    row[0].font = title_font
ws5.column_dimensions["A"].width = 96
ws5.column_dimensions["B"].width = 26
ws5["A1"].font = Font(name="Calibri", size=14, bold=True, color=MAROON)

wb.save(OUT)
print("wrote", OUT)
